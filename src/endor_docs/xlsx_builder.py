"""
EndorWorkbook — gera arquivos .xlsx aplicando a identidade Endor.

Uso típico:

    from endor_docs import EndorWorkbook
    from endor_docs.styles import BRAND

    wb = EndorWorkbook(title="Modelo Financeiro - Luz Imperial 2026")
    wb.add_cover(
        subtitle="PPP de Iluminação Pública - Vassouras",
        metadata={"Contrato": "27/2023", "Período": "Mar/2026"},
    )

    sheet = wb.add_sheet("Premissas")
    wb.write_header_row(sheet, row=1, headers=["Item", "Unidade", "Valor"])
    sheet.append(["Taxa de desconto", "% a.a.", 0.12])
    wb.format_percent(sheet, "C2:C2")

    wb.save("modelo.xlsx")

Convenções brasileiras de formatação numérica (BRL, kWh, MWh, percentuais com
vírgula decimal) são aplicadas via ``format_*`` helpers — o display depende
do locale do Excel do usuário, mas o número subjacente é correto.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from endor_docs import styles


# ---------------------------------------------------------------------------
# Formato numérico — Brasil
# ---------------------------------------------------------------------------

FORMAT_BRL = '_-"R$"\\ * #,##0.00_-;-"R$"\\ * #,##0.00_-;_-"R$"\\ * "-"??_-;_-@_-'
FORMAT_BRL_INT = '_-"R$"\\ * #,##0_-;-"R$"\\ * #,##0_-;_-"R$"\\ * "-"??_-;_-@_-'
FORMAT_INT = "#,##0"
FORMAT_DECIMAL = "#,##0.00"
FORMAT_PERCENT = "0.0%"
FORMAT_PERCENT_2 = "0.00%"
FORMAT_KWH = '#,##0\\ "kWh"'
FORMAT_MWH = '#,##0\\ "MWh"'
FORMAT_DATE_BR = "DD/MM/YYYY"


# ---------------------------------------------------------------------------
# Helpers de estilo
# ---------------------------------------------------------------------------


def _font(
    *,
    size: float = styles.SIZES.body,
    color: styles.Color = styles.NEUTRAL.gray_800,
    bold: bool = False,
    name: str | None = None,
) -> Font:
    return Font(
        name=name or styles.FONTS.document,
        size=size,
        bold=bold,
        color=color.argb,
    )


def _fill(color: styles.Color) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=color.argb, end_color=color.argb)


def _border_bottom(color: styles.Color = styles.NEUTRAL.gray_200, weight: str = "thin") -> Border:
    return Border(bottom=Side(border_style=weight, color=color.argb))


# ---------------------------------------------------------------------------
# EndorWorkbook
# ---------------------------------------------------------------------------


class EndorWorkbook:
    """Builder de planilhas .xlsx com identidade Endor aplicada."""

    def __init__(self, title: str = "", author: str = "Endor Energia"):
        self.wb = Workbook()
        # openpyxl cria uma sheet default chamada "Sheet" — vamos removê-la
        # depois que o usuário adicionar a primeira sheet de verdade.
        self._default_sheet_to_remove = self.wb.active.title
        self.wb.properties.title = title
        self.wb.properties.creator = author
        self._title = title

    # ------------------------------------------------------------------
    # Sheets primárias
    # ------------------------------------------------------------------

    def add_cover(
        self,
        subtitle: str = "",
        metadata: dict[str, str] | None = None,
    ) -> Worksheet:
        """Aba de capa com título grande + metadados em layout limpo."""
        ws = self.wb.create_sheet("Capa", 0)
        self._maybe_remove_default()

        # Largura das colunas
        ws.column_dimensions["A"].width = 2
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 30

        # Cor de fundo Endor cobrindo a área visível
        for row in range(1, 30):
            for col in range(1, 8):
                cell = ws.cell(row=row, column=col)
                cell.fill = _fill(styles.BRAND.endor)
        # Faixa decorativa terra no topo
        for col in range(1, 8):
            ws.cell(row=1, column=col).fill = _fill(styles.BRAND.terra)

        # Título
        ws.cell(row=4, column=2, value=self._title).font = _font(
            size=styles.SIZES.display,
            color=styles.NEUTRAL.white,
            bold=True,
        )
        ws.row_dimensions[4].height = 60

        if subtitle:
            ws.cell(row=7, column=2, value=subtitle).font = _font(
                size=styles.SIZES.h2,
                color=styles.NEUTRAL.white,
            )
            ws.row_dimensions[7].height = 24

        # Logo (imagem)
        if styles.LOGOS.fundo_azul_png.exists():
            img = OpenpyxlImage(str(styles.LOGOS.fundo_azul_png))
            img.height = 50
            img.width = 140  # proporcional ao logo
            ws.add_image(img, "B2")

        # Metadados em duas colunas (label + valor)
        if metadata:
            row = 12
            for key, val in metadata.items():
                label_cell = ws.cell(row=row, column=2, value=key.upper())
                label_cell.font = _font(
                    size=styles.SIZES.tiny,
                    color=styles.BRAND.terra,
                    bold=True,
                )
                val_cell = ws.cell(row=row, column=3, value=val)
                val_cell.font = _font(size=styles.SIZES.body, color=styles.NEUTRAL.white)
                row += 2

        ws.sheet_view.showGridLines = False
        return ws

    def add_sheet(self, name: str) -> Worksheet:
        """Cria uma nova aba e remove a default se ainda existir."""
        ws = self.wb.create_sheet(name)
        self._maybe_remove_default()
        ws.sheet_view.showGridLines = False
        return ws

    def _maybe_remove_default(self):
        if (
            self._default_sheet_to_remove
            and self._default_sheet_to_remove in self.wb.sheetnames
            and len(self.wb.sheetnames) > 1
        ):
            del self.wb[self._default_sheet_to_remove]
            self._default_sheet_to_remove = None

    # ------------------------------------------------------------------
    # Estilização de blocos
    # ------------------------------------------------------------------

    def write_header_row(
        self,
        ws: Worksheet,
        row: int,
        headers: Iterable[str],
        start_col: int = 1,
        color: styles.Color | None = None,
    ):
        """Escreve uma linha de cabeçalho com estilo padrão."""
        color = color or styles.BRAND.endor
        for i, header in enumerate(headers):
            cell = ws.cell(row=row, column=start_col + i, value=header)
            cell.font = _font(size=styles.SIZES.small, color=styles.NEUTRAL.white, bold=True)
            cell.fill = _fill(color)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 24

    def apply_zebra(
        self,
        ws: Worksheet,
        start_row: int,
        end_row: int,
        start_col: int = 1,
        end_col: int | None = None,
    ):
        """Aplica linhas alternadas (zebra) com cinza claro nas pares."""
        end_col = end_col or ws.max_column
        for row in range(start_row, end_row + 1):
            if (row - start_row) % 2 == 1:
                for col in range(start_col, end_col + 1):
                    cell = ws.cell(row=row, column=col)
                    cell.fill = _fill(styles.NEUTRAL.gray_50)

    def apply_table_borders(
        self,
        ws: Worksheet,
        start_row: int,
        end_row: int,
        start_col: int = 1,
        end_col: int | None = None,
    ):
        """Adiciona linha divisória inferior em cada linha (estilo Luz Imperial)."""
        end_col = end_col or ws.max_column
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                ws.cell(row=row, column=col).border = _border_bottom()

    # ------------------------------------------------------------------
    # Formatadores de coluna
    # ------------------------------------------------------------------

    def format_range(self, ws: Worksheet, cell_range: str, fmt: str):
        """Aplica um format code a um intervalo (ex: 'B2:B100')."""
        for row in ws[cell_range]:
            for cell in row:
                cell.number_format = fmt

    def format_brl(self, ws: Worksheet, cell_range: str, decimals: bool = True):
        self.format_range(ws, cell_range, FORMAT_BRL if decimals else FORMAT_BRL_INT)

    def format_int(self, ws: Worksheet, cell_range: str):
        self.format_range(ws, cell_range, FORMAT_INT)

    def format_decimal(self, ws: Worksheet, cell_range: str):
        self.format_range(ws, cell_range, FORMAT_DECIMAL)

    def format_percent(self, ws: Worksheet, cell_range: str, decimals: int = 1):
        self.format_range(ws, cell_range, FORMAT_PERCENT if decimals == 1 else FORMAT_PERCENT_2)

    def format_kwh(self, ws: Worksheet, cell_range: str):
        self.format_range(ws, cell_range, FORMAT_KWH)

    def format_mwh(self, ws: Worksheet, cell_range: str):
        self.format_range(ws, cell_range, FORMAT_MWH)

    def format_date(self, ws: Worksheet, cell_range: str):
        self.format_range(ws, cell_range, FORMAT_DATE_BR)

    # ------------------------------------------------------------------
    # Blocos visuais (KPIs)
    # ------------------------------------------------------------------

    def add_kpi_block(
        self,
        ws: Worksheet,
        start_row: int,
        kpis: list[dict],
        start_col: int = 2,
    ):
        """Insere uma linha de KPI cards horizontais.

        Cada KPI é dict com: ``label``, ``value``, ``subtext`` (opcional), ``color`` (opcional).
        """
        card_w = 3  # 3 colunas por card
        gap = 1
        col = start_col
        for kpi in kpis:
            color = kpi.get("color", styles.BRAND.terra)

            # Tag colorida na linha do topo do card
            for c in range(card_w):
                cell = ws.cell(row=start_row, column=col + c)
                cell.fill = _fill(color)
            ws.row_dimensions[start_row].height = 4

            # Background gray-50 nas linhas seguintes
            for r in range(start_row + 1, start_row + 5):
                for c in range(card_w):
                    cell = ws.cell(row=r, column=col + c)
                    cell.fill = _fill(styles.NEUTRAL.gray_50)

            # Label
            lbl_cell = ws.cell(row=start_row + 1, column=col, value=kpi["label"].upper())
            lbl_cell.font = _font(
                size=styles.SIZES.kpi_label,
                color=styles.NEUTRAL.gray_600,
                bold=True,
            )
            ws.merge_cells(
                start_row=start_row + 1,
                end_row=start_row + 1,
                start_column=col,
                end_column=col + card_w - 1,
            )

            # Value
            val_cell = ws.cell(row=start_row + 2, column=col, value=kpi["value"])
            val_cell.font = _font(
                size=styles.SIZES.kpi_value,
                color=styles.BRAND.endor,
                bold=True,
            )
            ws.row_dimensions[start_row + 2].height = 40
            ws.merge_cells(
                start_row=start_row + 2,
                end_row=start_row + 2,
                start_column=col,
                end_column=col + card_w - 1,
            )

            # Subtext
            if kpi.get("subtext"):
                sub_cell = ws.cell(row=start_row + 3, column=col, value=kpi["subtext"])
                sub_cell.font = _font(size=styles.SIZES.small, color=styles.NEUTRAL.gray_600)
                ws.merge_cells(
                    start_row=start_row + 3,
                    end_row=start_row + 3,
                    start_column=col,
                    end_column=col + card_w - 1,
                )

            col += card_w + gap

    # ------------------------------------------------------------------
    # Named ranges (modelos financeiros)
    # ------------------------------------------------------------------

    def define_name(self, name: str, sheet: Worksheet, cell_ref: str):
        """Define um named range (ex: 'taxa_desconto' -> Premissas!$C$2).

        Permite escrever fórmulas legíveis no modelo: ``=taxa_desconto * ...``.
        """
        full_ref = f"'{sheet.title}'!{cell_ref}"
        defined = DefinedName(name=name, attr_text=full_ref)
        self.wb.defined_names[name] = defined

    # ------------------------------------------------------------------
    # Save
    # ------------------------------------------------------------------

    def save(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self._maybe_remove_default()
        self.wb.save(str(path))
        return path
